"""Read the first sheet of .xlsx or a .csv into a header and rows of cells, without interpreting them."""

import csv
import io
from datetime import date, datetime

from openpyxl import load_workbook

MAX_BYTES = 5 * 1024 * 1024


class UnreadableFile(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def _cell(value: object) -> object:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableFile("bad_encoding")


def read_table(data: bytes, filename: str, max_rows: int) -> tuple[list[str], list[list[object]]]:
    """First non-empty row is the header. Raises UnreadableFile with a problem code."""
    if len(data) > MAX_BYTES:
        raise UnreadableFile("file_too_large")
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # zip or xml errors from a broken or renamed file
            raise UnreadableFile("bad_file") from e
        sheet = book.worksheets[0]
        raw = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        book.close()
    elif name.endswith((".csv", ".txt")):
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith(".xls"):
        raise UnreadableFile("old_excel")  # save as .xlsx or .csv
    else:
        raise UnreadableFile("bad_file")

    rows = [r for r in raw if any(str(c or "").strip() for c in r)]
    if not rows:
        raise UnreadableFile("empty_file")
    if len(rows) - 1 > max_rows:
        raise UnreadableFile("too_many_rows")
    width = max(len(r) for r in rows)
    padded = [list(r) + [None] * (width - len(r)) for r in rows]
    header = [str(c).strip() if c is not None else "" for c in padded[0]]
    return header, padded[1:]
